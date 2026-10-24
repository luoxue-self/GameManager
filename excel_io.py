"""Excel 导入导出：openpyxl 读写 xlsx + Pillow 处理图片。

导出：把本地图片作为浮动对象锚定到「游戏图片」列对应行，单元格里留空；
导入：解析表头（自动跳过标题行）、按图片 anchor 行号回填、复用 dedup 批量入库。

图片来源三种方式：
  A 内嵌  —— 图片插入到「游戏图片」列（最常用）
  B 文件名 —— 写 zelda.jpg，另附 zip 压缩包
  C 外链  —— 写 https://...，不下载，原样存库
"""
from __future__ import annotations

import io
import os
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from PIL import ImageOps

import dedup

HEADERS = ["游戏名", "游戏类型", "游戏编号", "游戏图片", "保存时间", "下载链接",
           "使用方法", "推荐指数", "本地状态", "游戏年份", "保存位置"]
FIELDS = ["name", "type", "code", "cover", "saved_at", "download_url",
          "how_to", "rating", "local_status", "year", "location"]
WIDTHS = [32, 12, 14, 22, 20, 42, 42, 10, 12, 10, 22]
IMG_H = 90       # 单元格内图片高度（像素）
IMG_MAX_W = 165  # 防止超宽封面撑破列宽
ROW_PT = 70      # 对应行高（磅）

HEADER_ALIAS = {
    "游戏名": "name", "名称": "name", "游戏名称": "name",
    "游戏类型": "type", "类型": "type",
    "游戏编号": "code", "编号": "code",
    "游戏图片": "cover", "图片": "cover", "封面": "cover",
    "保存时间": "saved_at", "时间": "saved_at",
    "下载链接": "download_url", "链接": "download_url",
    "使用方法": "how_to", "说明": "how_to",
    "推荐指数": "rating", "评分": "rating",
    "本地状态": "local_status", "状态": "local_status",
    "游戏年份": "year", "年份": "year",
    "保存位置": "location", "位置": "location",
}


# ---------------------------------------------------------------- 导出
def _safe_cover_path(img_dir: str, cover: str) -> str | None:
    """把 cover 解析为 img_dir 内的真实文件路径；越界或不存在返回 None。"""
    if not cover:
        return None
    base = os.path.realpath(img_dir)
    path = os.path.realpath(os.path.join(base, cover))
    if os.path.commonpath([base, path]) != base:
        return None
    return path if os.path.isfile(path) else None


def export_xlsx(rows, img_dir: str) -> io.BytesIO:
    """rows：可下标访问的映射（sqlite3.Row / dict）序列。返回 xlsx 二进制流。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "游戏库"
    ws.append(HEADERS)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    for r, row in enumerate(rows, 2):
        for c, f in enumerate(FIELDS, 1):
            if f == "cover":
                continue
            v = row[f]
            ws.cell(r, c, "" if v is None else v)
        cover = (row["cover"] or "").strip()
        if cover and not cover.startswith(("http://", "https://")):
            path = _safe_cover_path(img_dir, cover)
            if path:
                try:
                    img = XLImage(path)
                    ratio = IMG_H / img.height if img.height else 1
                    img.height = IMG_H
                    img.width = min(IMG_MAX_W, max(1, int(img.width * ratio)))
                    ws.add_image(img, f"D{r}")
                    ws.row_dimensions[r].height = ROW_PT
                    ws.cell(r, 4).alignment = Alignment(
                        vertical="center", horizontal="center")
                    continue
                except Exception:
                    pass
        ws.cell(r, 4, cover)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- 图片
def _img_bytes(img) -> bytes | None:
    """兼容不同 openpyxl 版本读取内嵌图片原始字节。"""
    try:
        return img._data()
    except Exception:
        pass
    ref = getattr(img, "ref", None)
    if hasattr(ref, "getvalue"):
        return ref.getvalue()
    if hasattr(ref, "read"):
        pos = ref.tell()
        ref.seek(0)
        data = ref.read()
        ref.seek(pos)
        return data
    return None


def save_image_bytes(data: bytes, img_dir: str,
                     max_side: int | None = None) -> str:
    """保存图片，返回文件名。

    max_side=None：原样保存（导入用，保留原始质量）。
    max_side=N   ：等比缩到最长边 N 并重编码（上传用，避免大图拖慢列表）。
    """
    os.makedirs(img_dir, exist_ok=True)
    if max_side is None:
        im = PILImage.open(io.BytesIO(data))
        fmt = (im.format or "PNG").lower()
        ext = "jpg" if fmt in ("jpeg", "jpg") else fmt
        if ext not in ("jpg", "png", "gif", "webp", "bmp"):
            # 浏览器不认识的格式统一转成 PNG
            reencoded = io.BytesIO()
            im.convert("RGBA").save(reencoded, "PNG")
            data, ext = reencoded.getvalue(), "png"
        fn = f"{uuid.uuid4().hex}.{ext}"
        with open(os.path.join(img_dir, fn), "wb") as f:
            f.write(data)
        return fn

    im = PILImage.open(io.BytesIO(data))
    im = ImageOps.exif_transpose(im)
    im.thumbnail((max_side, max_side), PILImage.LANCZOS)
    has_alpha = im.mode in ("RGBA", "LA") or (
        im.mode == "P" and "transparency" in im.info)
    fn = f"{uuid.uuid4().hex}." + ("png" if has_alpha else "jpg")
    path = os.path.join(img_dir, fn)
    if has_alpha:
        im.convert("RGBA").save(path, "PNG", optimize=True)
    else:
        im.convert("RGB").save(path, "JPEG", quality=88,
                               optimize=True, progressive=True)
    return fn


def _num(v, default=0.0) -> float:
    try:
        return float(v)
    except Exception:
        return default


def _int(v):
    try:
        return int(float(v))
    except Exception:
        return None


# ---------------------------------------------------------------- 导入
def _find_header(ws, max_scan: int = 10) -> tuple[int, dict[int, str]]:
    """在前 10 行里找表头，返回 (表头所在行号, {列号: 字段名})。"""
    for row in range(1, min(ws.max_row, max_scan) + 1):
        header = {}
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row, col).value
            if isinstance(v, str):
                v = v.strip()
                if v in HEADER_ALIAS:
                    header[col] = HEADER_ALIAS[v]
        if "name" in header.values():
            return row, header
    return 0, {}


def _match_existing(existing, name: str, norm: str, threshold: float):
    """在库里找最合适的已有记录。

    优先级：完全同名 > 归一化同名 > 相似度最高且 ≥ 阈值。
    （同归一化名的记录可能有多条，比如被 force 保存过的近似条目。）
    """
    exact_norm = None
    best = None
    for gid, gname, gnorm in existing:
        if gname == name:
            return gid, gname
        if exact_norm is None and gnorm == norm:
            exact_norm = (gid, gname)
        score = dedup.similarity(norm, gnorm)
        if best is None or score > best[2]:
            best = (gid, gname, score)
    if exact_norm is not None:
        return exact_norm
    if best is not None and best[2] >= threshold:
        return best[0], best[1]
    return None, None


def _inline_images(ws) -> dict[int, bytes]:
    """{Excel 行号(1 起): 图片字节}"""
    out: dict[int, bytes] = {}
    for img in getattr(ws, "_images", []):
        try:
            anchor = getattr(img, "anchor", None)
            frm = getattr(anchor, "_from", None)
            if frm is None:
                continue
            r = frm.row + 1
            data = _img_bytes(img)
            if data:
                out.setdefault(r, data)
        except Exception:
            continue
    return out


def _remove_image(img_dir: str, filename: str) -> None:
    if filename and not filename.startswith(("http://", "https://")):
        try:
            os.remove(os.path.join(img_dir, filename))
        except OSError:
            pass


def import_xlsx(xlsx_bytes: bytes, con, img_dir: str,
                zip_images: dict | None = None,
                dry_run: bool = True,
                dup_threshold: float = 0.85,
                on_dup: str = "skip") -> dict:
    """把 xlsx 导入 games 表。

    dry_run=True  只预检，不写库、不落图。
    on_dup="skip" 重复条目跳过（默认）；"update" 覆盖已有条目（保留旧封面）。
    """
    on_dup = "update" if on_dup == "update" else "skip"
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws, header_row, header = None, 0, {}
    for sheet in wb.worksheets:  # 优先找含「游戏名」表头的工作表
        hr, hd = _find_header(sheet)
        if "name" in hd.values():
            ws, header_row, header = sheet, hr, hd
            break
    if ws is None:
        return {"ok": False, "error": "表头缺少『游戏名』列（未找到含该列的工作表）"}

    inline_imgs = _inline_images(ws)
    zip_images = {k.lower(): v for k, v in (zip_images or {}).items()}
    os.makedirs(img_dir, exist_ok=True)

    existing = [(r["id"], r["name"], r["name_norm"])
                for r in con.execute("SELECT id,name,name_norm FROM games")]
    stats = {"total": 0, "inserted": 0, "updated": 0, "skipped_dup": 0,
             "skipped_invalid": 0, "images": 0, "details": []}
    batch: list[dict] = []
    touched: set[int] = set()

    for r in range(header_row + 1, ws.max_row + 1):
        data = {}
        for col, field in header.items():
            v = ws.cell(r, col).value
            data[field] = "" if v is None else str(v).strip()
        if not any(data.values()):
            continue
        stats["total"] += 1
        name = data.get("name", "").strip()
        if not name:
            stats["skipped_invalid"] += 1
            stats["details"].append({"row": r, "name": "", "status": "缺游戏名"})
            continue

        norm = dedup.normalize(name)
        hit_id, hit_name = _match_existing(existing, name, norm, dup_threshold)
        if hit_id is None:
            for b in batch:
                if b["name_norm"] == norm or \
                        dedup.similarity(norm, b["name_norm"]) >= dup_threshold:
                    hit_name = b["name"] + "（本表内）"
                    break

        # 不能安全更新：本表内重复，或本次导入已经更新过同一条，或策略为跳过
        if hit_name and (hit_id is None or on_dup != "update" or hit_id in touched):
            stats["skipped_dup"] += 1
            status = f"疑似重复 → {hit_name}"
            if hit_id in touched:
                status = "本表内重复出现，已跳过"
            stats["details"].append({"row": r, "name": name, "status": status})
            continue

        cover = ""
        saved_image = ""
        ref = data.get("cover", "")
        if r in inline_imgs:
            cover = "（内嵌）"
            if not dry_run:
                cover = save_image_bytes(inline_imgs[r], img_dir)
                saved_image = cover
            stats["images"] += 1
        elif ref and ref.lower() in zip_images:
            cover = "（zip）"
            if not dry_run:
                cover = save_image_bytes(zip_images[ref.lower()], img_dir)
                saved_image = cover
            stats["images"] += 1
        elif ref.startswith(("http://", "https://")):
            cover = ref

        py_full, py_ini = dedup.pinyin(name)
        rec = {
            "name": name, "name_norm": norm,
            "name_pinyin": py_full, "name_initials": py_ini,
            "type": data.get("type", ""), "code": data.get("code") or None,
            "cover": cover, "saved_at": data.get("saved_at") or None,
            "download_url": data.get("download_url", ""),
            "how_to": data.get("how_to", ""),
            "rating": _num(data.get("rating"), 0),
            "local_status": data.get("local_status") or "未下载",
            "year": _int(data.get("year")),
            "location": data.get("location", ""),
        }

        # ---------- 覆盖已有条目 ----------
        if hit_id is not None and on_dup == "update":
            touched.add(hit_id)
            if dry_run:
                stats["updated"] += 1
                stats["details"].append({"row": r, "name": name,
                                         "status": "将更新（预检）"})
                continue
            old = con.execute("SELECT cover, saved_at FROM games WHERE id=?",
                              (hit_id,)).fetchone()
            if old is not None:
                rec["cover"] = rec["cover"] or (old["cover"] or "")
                rec["saved_at"] = rec["saved_at"] or old["saved_at"]
                try:
                    con.execute(
                        """UPDATE games SET name=?,name_norm=?,name_pinyin=?,
                           name_initials=?,type=?,code=?,cover=?,saved_at=?,
                           download_url=?,how_to=?,rating=?,local_status=?,
                           year=?,location=? WHERE id=?""",
                        (rec["name"], rec["name_norm"], rec["name_pinyin"],
                         rec["name_initials"], rec["type"], rec["code"],
                         rec["cover"], rec["saved_at"], rec["download_url"],
                         rec["how_to"], rec["rating"], rec["local_status"],
                         rec["year"], rec["location"], hit_id))
                    stats["updated"] += 1
                    stats["details"].append({"row": r, "name": name,
                                             "status": "已更新"})
                except Exception as e:
                    stats["skipped_invalid"] += 1
                    msg = ("编号已存在" if "UNIQUE" in str(e).upper()
                           else f"更新失败:{e}")
                    stats["details"].append({"row": r, "name": name,
                                             "status": msg})
                    _remove_image(img_dir, saved_image)
                continue

        # ---------- 新增 ----------
        batch.append(rec)
        if dry_run:
            stats["inserted"] += 1
            stats["details"].append({"row": r, "name": name,
                                     "status": "可导入（预检）"})
            continue
        try:
            con.execute(
                """INSERT INTO games(name,name_norm,name_pinyin,name_initials,
                   type,code,cover,saved_at,download_url,how_to,rating,
                   local_status,year,location)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (rec["name"], rec["name_norm"], rec["name_pinyin"],
                 rec["name_initials"], rec["type"], rec["code"], rec["cover"],
                 rec["saved_at"], rec["download_url"], rec["how_to"],
                 rec["rating"], rec["local_status"], rec["year"],
                 rec["location"]))
            stats["inserted"] += 1
            stats["details"].append({"row": r, "name": name, "status": "已导入"})
        except Exception as e:
            stats["skipped_invalid"] += 1
            msg = "编号已存在" if "UNIQUE" in str(e).upper() else f"失败:{e}"
            stats["details"].append({"row": r, "name": name, "status": msg})
            _remove_image(img_dir, saved_image)

    if not dry_run:
        con.commit()
    stats["ok"] = True
    stats["dry_run"] = dry_run
    stats["on_dup"] = on_dup
    return stats
