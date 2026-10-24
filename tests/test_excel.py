import io
import os

from openpyxl import load_workbook
from PIL import Image

import db
import dedup
import excel_io


def _cover(color) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (600, 800), color).save(buf, "JPEG", quality=90)
    return buf.getvalue()


def _clear():
    con = db.connect()
    con.execute("DELETE FROM games")
    con.commit()
    return con


def test_export_import_roundtrip():
    con = _clear()
    samples = [
        ("塞尔达传说：王国之泪", (200, 60, 60), "RPG", "T1"),
        ("博德之门3", (60, 160, 90), "RPG", "T2"),
        ("星露谷物语", (70, 110, 220), "模拟", ""),
    ]
    for i, (name, color, gtype, code) in enumerate(samples):
        fn = excel_io.save_image_bytes(_cover(color), db.IMG_DIR)
        con.execute(
            """INSERT INTO games(name,name_norm,type,code,cover,saved_at,
               download_url,how_to,rating,local_status,year,location)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (name, dedup.normalize(name), gtype, code or None, fn,
             "2024-05-01 10:00", "https://example.com/x", "解压即玩", 4.5,
             "已下载", 2023, "D:/Games"))
    con.commit()

    rows = [dict(r) for r in con.execute("SELECT * FROM games ORDER BY id")]
    data = excel_io.export_xlsx(rows, db.IMG_DIR).getvalue()
    con.close()

    # 表头 + 三张内嵌图片
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.cell(1, 1).value == "游戏名"
    assert ws.cell(1, 4).value == "游戏图片"
    assert len(ws._images) == 3

    # 预检：不写库、不落图
    con = _clear()
    before = con.execute("SELECT COUNT(*) c FROM games").fetchone()["c"]
    assert before == 0
    dry = excel_io.import_xlsx(data, con, db.IMG_DIR, dry_run=True)
    assert dry["ok"] and dry["total"] == 3
    assert dry["inserted"] == 3 and dry["images"] == 3 and dry["skipped_dup"] == 0
    assert con.execute("SELECT COUNT(*) c FROM games").fetchone()["c"] == 0

    # 正式导入
    res = excel_io.import_xlsx(data, con, db.IMG_DIR, dry_run=False)
    assert res["inserted"] == 3 and res["images"] == 3 and res["skipped_dup"] == 0
    assert con.execute("SELECT COUNT(*) c FROM games").fetchone()["c"] == 3
    for r in con.execute("SELECT * FROM games"):
        assert r["cover"] and not r["cover"].startswith("http")
        assert os.path.isfile(os.path.join(db.IMG_DIR, r["cover"]))
        assert r["rating"] == 4.5 and r["year"] == 2023

    # 再导一次 → 全部判重
    again = excel_io.import_xlsx(data, con, db.IMG_DIR, dry_run=False)
    assert again["inserted"] == 0 and again["skipped_dup"] == 3
    con.close()


def test_import_rejects_sheet_without_name_header():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "类型"])
    ws.append(["X1", "RPG"])
    buf = io.BytesIO()
    wb.save(buf)

    con = _clear()
    res = excel_io.import_xlsx(buf.getvalue(), con, db.IMG_DIR, dry_run=True)
    con.close()
    assert res["ok"] is False
    assert "游戏名" in res["error"]


def test_import_uses_sheet_with_name_header():
    from openpyxl import Workbook

    wb = Workbook()
    cover_sheet = wb.active
    cover_sheet.title = "说明"
    cover_sheet["A1"] = "这是一个说明页，没有表头"
    ws = wb.create_sheet("游戏库")
    ws.append(["我的游戏库"])          # 标题行，应被跳过
    ws.append(excel_io.HEADERS)
    ws.append(["分表游戏", "动作", "S1", "", "", "", "", "4.5", "未下载", "2020", ""])
    buf = io.BytesIO()
    wb.save(buf)

    con = _clear()
    res = excel_io.import_xlsx(buf.getvalue(), con, db.IMG_DIR, dry_run=False)
    assert res["ok"] and res["inserted"] == 1, res
    row = con.execute("SELECT * FROM games").fetchone()
    assert row["name"] == "分表游戏"
    assert row["year"] == 2020 and row["rating"] == 4.5
    con.execute("DELETE FROM games")
    con.commit()
    con.close()


def test_import_update_existing_mode():
    from openpyxl import Workbook

    con = _clear()
    py_full, py_ini = dedup.pinyin("更新目标")
    con.execute(
        """INSERT INTO games(name,name_norm,name_pinyin,name_initials,type,
           code,cover,saved_at,download_url,how_to,rating,local_status,year,location)
           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        ("更新目标", dedup.normalize("更新目标"), py_full, py_ini, "旧类型",
         "U1", "keep.jpg", "2024-01-01 00:00", "", "旧说明", 3,
         "未下载", 2020, ""))
    con.commit()

    wb = Workbook()
    ws = wb.active
    ws.append(excel_io.HEADERS)
    ws.append(["更新目标", "新类型", "U1", "", "2024-02-02 10:00", "",
               "新说明", "4.5", "已通关", "2021", "D:/new"])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    # 默认跳过
    skip = excel_io.import_xlsx(data, con, db.IMG_DIR, dry_run=True)
    assert skip["skipped_dup"] == 1 and skip["inserted"] == 0
    assert skip["updated"] == 0

    # 预检更新
    dry = excel_io.import_xlsx(data, con, db.IMG_DIR, dry_run=True,
                               on_dup="update")
    assert dry["updated"] == 1 and dry["inserted"] == 0

    # 正式更新
    res = excel_io.import_xlsx(data, con, db.IMG_DIR, dry_run=False,
                               on_dup="update")
    assert res["updated"] == 1 and res["inserted"] == 0
    assert res["on_dup"] == "update"

    row = con.execute("SELECT * FROM games").fetchone()
    assert row["type"] == "新类型" and row["how_to"] == "新说明"
    assert row["local_status"] == "已通关" and row["year"] == 2021
    assert row["rating"] == 4.5
    assert row["cover"] == "keep.jpg"          # 新表没有图片，保留旧封面
    assert row["name_pinyin"] == py_full
    assert row["name_initials"] == py_ini

    con.execute("DELETE FROM games")
    con.commit()
    con.close()


def test_import_matches_zip_image_by_filename():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(excel_io.HEADERS)
    ws.append(["只有文件名的游戏", "动作", "Z1", "zelda.jpg", "", "", "", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)

    con = _clear()
    res = excel_io.import_xlsx(buf.getvalue(), con, db.IMG_DIR, dry_run=False,
                               zip_images={"zelda.jpg": _cover((10, 10, 10))})
    assert res["inserted"] == 1 and res["images"] == 1
    row = con.execute("SELECT * FROM games").fetchone()
    assert row["cover"] and os.path.isfile(os.path.join(db.IMG_DIR, row["cover"]))
    con.execute("DELETE FROM games")
    con.commit()
    con.close()
